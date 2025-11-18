"""
Excel Data Extraction Functions

Provides specialized extraction functions for different data patterns
including structured tables, key-value pairs, time-series, and metadata.
"""

import re
from typing import Any, Dict, List, Optional, Tuple, Union

import pandas as pd
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


class ExcelExtractor:
    """
    Extractor for various Excel data patterns.

    Handles:
    - Structured tables with headers
    - Key-value pairs (equipment specs)
    - Time-series data
    - Metadata headers
    - Multi-level headers
    - Cell comments and notes
    - Named ranges
    """

    def __init__(self):
        """Initialize extractor."""
        pass

    def extract_range(
        self,
        workbook: Workbook,
        sheet_name: str,
        cell_range: str,
        header_row: Optional[int] = 0
    ) -> pd.DataFrame:
        """
        Extract data from a specific cell range.

        Args:
            workbook: openpyxl Workbook object
            sheet_name: Name of the sheet
            cell_range: Excel range notation (e.g., "A1:D100")
            header_row: Row index for column headers (0-indexed, None for no header)

        Returns:
            DataFrame with extracted data
        """
        if workbook is None:
            raise ValueError("Workbook is None - cannot extract range")

        sheet = workbook[sheet_name]

        # Parse range
        start_cell, end_cell = cell_range.split(':')
        data = []

        # Get cell range values
        for row in sheet[cell_range]:
            row_data = [cell.value for cell in row]
            data.append(row_data)

        if not data:
            return pd.DataFrame()

        # Convert to DataFrame
        if header_row is not None and len(data) > header_row:
            headers = data[header_row]
            df = pd.DataFrame(data[header_row + 1:], columns=headers)
        else:
            df = pd.DataFrame(data)

        return df

    def extract_key_value_pairs(
        self,
        workbook: Workbook,
        sheet_name: str,
        search_range: str = "A1:B50"
    ) -> Dict[str, Any]:
        """
        Extract key-value pairs from equipment specs or metadata sections.

        Searches for patterns like:
        - Equipment ID: XYZ-123
        - Test Date: 2024-01-15
        - Operator: John Doe

        Handles both:
        - Single cell format: "Key: Value"
        - Two column format: | Key | Value |

        Args:
            workbook: openpyxl Workbook object
            sheet_name: Name of the sheet
            search_range: Range to search for key-value pairs

        Returns:
            Dictionary of extracted key-value pairs
        """
        metadata = {}

        if workbook is None:
            return metadata

        sheet = workbook[sheet_name]

        # Get cell range
        for row in sheet[search_range]:
            # Try two column format first
            if len(row) >= 2:
                key_cell = row[0]
                value_cell = row[1]

                if key_cell.value and value_cell.value:
                    key = str(key_cell.value).strip()
                    value = value_cell.value

                    # Clean up key (remove colons, etc.)
                    key = key.rstrip(':').strip()

                    metadata[key] = value
                    continue  # Skip single cell format check for this row

            # Check first cell for single cell format "Key: Value"
            # Only if two-column format didn't match
            if len(row) >= 1 and row[0].value:
                cell_value = str(row[0].value)

                # Look for colon separator
                if ':' in cell_value:
                    parts = cell_value.split(':', 1)
                    if len(parts) == 2:
                        key = parts[0].strip()
                        value = parts[1].strip()
                        # Only add if value is not empty
                        if value:
                            metadata[key] = value

        return metadata

    def extract_time_series(
        self,
        workbook: Workbook,
        sheet_name: str,
        time_column: str,
        value_columns: List[str],
        data_range: Optional[str] = None
    ) -> pd.DataFrame:
        """
        Extract time-series data with timestamp index.

        Args:
            workbook: openpyxl Workbook object
            sheet_name: Name of the sheet
            time_column: Name of the time/timestamp column
            value_columns: List of value column names
            data_range: Excel range notation (optional)

        Returns:
            DataFrame with time-series data
        """
        if data_range:
            df = self.extract_range(workbook, sheet_name, data_range)
        else:
            # Read entire sheet
            df = pd.read_excel(
                workbook,
                sheet_name=sheet_name
            )

        # Validate columns exist
        missing_cols = [col for col in [time_column] + value_columns if col not in df.columns]
        if missing_cols:
            raise ValueError(f"Missing columns: {missing_cols}")

        # Convert time column to datetime
        df[time_column] = pd.to_datetime(df[time_column], errors='coerce')

        # Set time as index
        df = df.set_index(time_column)

        # Select only requested columns
        df = df[value_columns]

        return df

    def extract_multi_level_headers(
        self,
        workbook: Workbook,
        sheet_name: str,
        header_rows: List[int],
        data_start_row: int
    ) -> pd.DataFrame:
        """
        Extract data with multi-level (merged) column headers.

        Example:
            Row 1: | Test Conditions |  Measurements  |
            Row 2: |  Temp  |   RH   |   V   |   I   |

        Args:
            workbook: openpyxl Workbook object
            sheet_name: Name of the sheet
            header_rows: List of row indices for headers (0-indexed)
            data_start_row: Row index where data starts

        Returns:
            DataFrame with multi-level column names
        """
        sheet = workbook[sheet_name]

        # Extract header rows
        headers = []
        for row_idx in header_rows:
            row = list(sheet.iter_rows(
                min_row=row_idx + 1,
                max_row=row_idx + 1,
                values_only=True
            ))[0]
            headers.append(row)

        # Combine multi-level headers
        combined_headers = []
        for col_idx in range(len(headers[0])):
            header_parts = []
            for row in headers:
                if col_idx < len(row) and row[col_idx]:
                    header_parts.append(str(row[col_idx]))

            # Join header parts
            combined_header = ' - '.join(header_parts) if header_parts else f'Column_{col_idx}'
            combined_headers.append(combined_header)

        # Extract data rows
        data = []
        for row in sheet.iter_rows(
            min_row=data_start_row + 1,
            values_only=True
        ):
            if any(cell is not None for cell in row):
                data.append(list(row))

        # Create DataFrame
        df = pd.DataFrame(data, columns=combined_headers)

        return df

    def extract_cell_comments(
        self,
        workbook: Workbook,
        sheet_name: str
    ) -> Dict[str, str]:
        """
        Extract cell comments and notes from a sheet.

        Args:
            workbook: openpyxl Workbook object
            sheet_name: Name of the sheet

        Returns:
            Dictionary mapping cell addresses to comment text
        """
        comments = {}
        sheet = workbook[sheet_name]

        for row in sheet.iter_rows():
            for cell in row:
                if cell.comment:
                    cell_address = f"{get_column_letter(cell.column)}{cell.row}"
                    comments[cell_address] = cell.comment.text

        return comments

    def extract_named_ranges(
        self,
        workbook: Workbook
    ) -> Dict[str, str]:
        """
        Extract named ranges from workbook.

        Args:
            workbook: openpyxl Workbook object

        Returns:
            Dictionary mapping named range names to their addresses
        """
        named_ranges = {}

        for name, definition in workbook.defined_names.items():
            # Get the range reference
            if definition.value:
                named_ranges[name] = definition.value

        return named_ranges

    def extract_conditional_formatting(
        self,
        workbook: Workbook,
        sheet_name: str
    ) -> List[Dict[str, Any]]:
        """
        Extract conditional formatting rules from a sheet.

        Args:
            workbook: openpyxl Workbook object
            sheet_name: Name of the sheet

        Returns:
            List of conditional formatting rules
        """
        rules = []
        sheet = workbook[sheet_name]

        for cf_range, cf_rules in sheet.conditional_formatting._cf_rules.items():
            for rule in cf_rules:
                rules.append({
                    'range': str(cf_range),
                    'type': rule.type,
                    'priority': rule.priority
                })

        return rules

    def extract_cell_styles(
        self,
        workbook: Workbook,
        sheet_name: str,
        cell_range: Optional[str] = None
    ) -> pd.DataFrame:
        """
        Extract cell formatting and styles (colors, fonts, etc.).

        Args:
            workbook: openpyxl Workbook object
            sheet_name: Name of the sheet
            cell_range: Excel range notation (None for entire sheet)

        Returns:
            DataFrame with cell addresses and their styles
        """
        sheet = workbook[sheet_name]

        style_data = []

        if cell_range:
            cells = sheet[cell_range]
        else:
            cells = sheet.iter_rows()

        for row in cells:
            for cell in row:
                if cell.value is not None:
                    cell_address = f"{get_column_letter(cell.column)}{cell.row}"

                    style_info = {
                        'cell': cell_address,
                        'value': cell.value,
                        'font_name': cell.font.name if cell.font else None,
                        'font_size': cell.font.size if cell.font else None,
                        'font_bold': cell.font.bold if cell.font else False,
                        'font_color': cell.font.color.rgb if (cell.font and cell.font.color) else None,
                        'fill_color': cell.fill.start_color.rgb if (cell.fill and hasattr(cell.fill, 'start_color')) else None,
                        'number_format': cell.number_format
                    }

                    style_data.append(style_info)

        return pd.DataFrame(style_data)

    def extract_merged_cells(
        self,
        workbook: Workbook,
        sheet_name: str
    ) -> List[str]:
        """
        Extract merged cell ranges from a sheet.

        Args:
            workbook: openpyxl Workbook object
            sheet_name: Name of the sheet

        Returns:
            List of merged cell ranges
        """
        sheet = workbook[sheet_name]
        merged_ranges = [str(merged_range) for merged_range in sheet.merged_cells.ranges]
        return merged_ranges

    def extract_images(
        self,
        workbook: Workbook,
        sheet_name: str
    ) -> List[Dict[str, Any]]:
        """
        Extract embedded images metadata from a sheet.

        Args:
            workbook: openpyxl Workbook object
            sheet_name: Name of the sheet

        Returns:
            List of image metadata dictionaries
        """
        sheet = workbook[sheet_name]
        images = []

        if hasattr(sheet, '_images'):
            for idx, img in enumerate(sheet._images):
                image_info = {
                    'index': idx,
                    'anchor': str(img.anchor) if hasattr(img, 'anchor') else None,
                    'format': img.format if hasattr(img, 'format') else None,
                }
                images.append(image_info)

        return images

    def extract_charts(
        self,
        workbook: Workbook,
        sheet_name: str
    ) -> List[Dict[str, Any]]:
        """
        Extract embedded charts metadata from a sheet.

        Args:
            workbook: openpyxl Workbook object
            sheet_name: Name of the sheet

        Returns:
            List of chart metadata dictionaries
        """
        sheet = workbook[sheet_name]
        charts = []

        if hasattr(sheet, '_charts'):
            for idx, chart in enumerate(sheet._charts):
                chart_info = {
                    'index': idx,
                    'type': chart.__class__.__name__,
                    'anchor': str(chart.anchor) if hasattr(chart, 'anchor') else None,
                }
                charts.append(chart_info)

        return charts

    def detect_table_boundaries(
        self,
        workbook: Workbook,
        sheet_name: str,
        start_cell: str = "A1"
    ) -> Tuple[str, str]:
        """
        Automatically detect the boundaries of a data table.

        Args:
            workbook: openpyxl Workbook object
            sheet_name: Name of the sheet
            start_cell: Starting cell to search from

        Returns:
            Tuple of (start_cell, end_cell) defining the table range
        """
        sheet = workbook[sheet_name]

        # Parse start cell
        start_col = sheet[start_cell].column
        start_row = sheet[start_cell].row

        # Find last column with data
        max_col = start_col
        for cell in sheet[start_row]:
            if cell.value is not None:
                max_col = max(max_col, cell.column)

        # Find last row with data
        max_row = start_row
        for row_idx, row in enumerate(sheet.iter_rows(min_row=start_row), start=start_row):
            if any(cell.value is not None for cell in row):
                max_row = row_idx

        # Build range
        start_cell_addr = f"{get_column_letter(start_col)}{start_row}"
        end_cell_addr = f"{get_column_letter(max_col)}{max_row}"

        return start_cell_addr, end_cell_addr

    def extract_formulas(
        self,
        workbook: Workbook,
        sheet_name: str,
        cell_range: Optional[str] = None
    ) -> Dict[str, str]:
        """
        Extract cell formulas from a sheet.

        Note: This requires loading workbook with data_only=False

        Args:
            workbook: openpyxl Workbook object
            sheet_name: Name of the sheet
            cell_range: Excel range notation (None for entire sheet)

        Returns:
            Dictionary mapping cell addresses to formulas
        """
        # Note: This won't work if workbook loaded with data_only=True
        # Would need to reload the workbook
        from openpyxl import load_workbook

        # Reload without data_only to get formulas
        formula_wb = load_workbook(
            filename=workbook._archive.filename,
            data_only=False
        )

        sheet = formula_wb[sheet_name]
        formulas = {}

        if cell_range:
            cells = sheet[cell_range]
        else:
            cells = sheet.iter_rows()

        for row in cells:
            for cell in row:
                if cell.value and isinstance(cell.value, str) and cell.value.startswith('='):
                    cell_address = f"{get_column_letter(cell.column)}{cell.row}"
                    formulas[cell_address] = cell.value

        formula_wb.close()

        return formulas
