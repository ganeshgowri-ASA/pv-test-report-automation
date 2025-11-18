"""
Unit Tests for Excel Data Ingestion Engine

Comprehensive test suite for Excel parser, validators, extractors, and transformers.
"""

import os
import tempfile
import unittest
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook

from excel_extractors import ExcelExtractor
from excel_parser import (
    CorruptedFileError,
    ExcelIngestion,
    ExcelIngestionResult,
    FileNotFoundError,
    InvalidSheetError,
)
from excel_transformers import ExcelTransformer
from excel_validators import ExcelValidator


class TestExcelValidator(unittest.TestCase):
    """Test cases for ExcelValidator class."""

    def setUp(self):
        """Set up test fixtures."""
        self.validator = ExcelValidator()

    def test_validate_missing_values(self):
        """Test missing value detection."""
        df = pd.DataFrame({
            'A': [1, 2, None, 4, 5],
            'B': [1, 2, 3, 4, 5],
            'C': [None, None, None, None, None]
        })

        result = self.validator.validate_missing_values(df)

        self.assertIn('warnings', result)
        self.assertIn('errors', result)
        # Column C has >50% missing, should be in errors
        self.assertTrue(any('C' in err for err in result['errors']))

    def test_validate_data_types(self):
        """Test data type validation."""
        df = pd.DataFrame({
            'A': ['1', '2', '3', '4', '5'],  # Should be numeric
            'B': ['a', 'b', 'c', 'd', 'e']   # Should stay text
        })

        result = self.validator.validate_data_types(df)

        # Column A should trigger warning (mostly numeric)
        self.assertTrue(any('A' in warn for warn in result['warnings']))

    def test_validate_duplicates(self):
        """Test duplicate detection."""
        df = pd.DataFrame({
            'A': [1, 2, 3, 2, 1],
            'B': ['a', 'b', 'c', 'b', 'a']
        })

        result = self.validator.validate_duplicates(df)

        self.assertTrue(len(result['warnings']) > 0)
        self.assertIn('duplicate', result['warnings'][0].lower())

    def test_detect_outliers_zscore(self):
        """Test outlier detection using Z-score method."""
        np.random.seed(42)
        data = np.random.normal(100, 10, 100)
        # Add outliers
        data = np.append(data, [200, 250, 0, -50])

        df = pd.DataFrame({'value': data})

        result = self.validator.detect_outliers(df, method='zscore')

        self.assertTrue(len(result['warnings']) > 0)
        self.assertIn('outliers', result['warnings'][0].lower())

    def test_detect_outliers_iqr(self):
        """Test outlier detection using IQR method."""
        np.random.seed(42)
        data = np.random.normal(100, 10, 100)
        data = np.append(data, [200, 250])

        df = pd.DataFrame({'value': data})

        result = self.validator.detect_outliers(df, method='iqr')

        self.assertTrue(len(result['warnings']) > 0)

    def test_validate_unit_consistency(self):
        """Test unit consistency validation."""
        # Values in mV (should be V)
        df = pd.DataFrame({
            'Voltage (V)': [37500, 38000, 37800]  # Actually in mV
        })

        result = self.validator.validate_unit_consistency(
            df,
            'Voltage (V)',
            expected_range=(30, 45),
            unit_multipliers={'mV': 0.001, 'V': 1.0}
        )

        self.assertIsNotNone(result['suggested_conversion'])
        self.assertEqual(result['suggested_conversion'][0], 'mV')

    def test_validate_timestamp_format(self):
        """Test timestamp format validation."""
        df = pd.DataFrame({
            'timestamp': ['2024-01-01', '2024-01-02', 'invalid', '2024-01-04']
        })

        result = self.validator.validate_timestamp_format(df, 'timestamp')

        self.assertTrue(len(result['warnings']) > 0)


class TestExcelExtractor(unittest.TestCase):
    """Test cases for ExcelExtractor class."""

    def setUp(self):
        """Set up test fixtures."""
        self.extractor = ExcelExtractor()
        self.temp_dir = tempfile.mkdtemp()

    def tearDown(self):
        """Clean up temporary files."""
        import shutil
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def create_test_workbook(self):
        """Create a test workbook."""
        wb = Workbook()
        ws = wb.active
        ws.title = "TestSheet"

        # Add metadata
        ws['A1'] = 'Equipment ID:'
        ws['B1'] = 'TEST-123'
        ws['A2'] = 'Test Date:'
        ws['B2'] = '2024-01-15'

        # Add data table
        ws['A5'] = 'Time'
        ws['B5'] = 'Value'
        ws['A6'] = 0
        ws['B6'] = 100
        ws['A7'] = 1
        ws['B7'] = 150

        return wb

    def test_extract_range(self):
        """Test extracting data from a specific range."""
        wb = self.create_test_workbook()

        df = self.extractor.extract_range(wb, 'TestSheet', 'A5:B7', header_row=0)

        self.assertEqual(len(df), 2)
        self.assertIn('Time', df.columns)
        self.assertIn('Value', df.columns)

    def test_extract_key_value_pairs(self):
        """Test extracting key-value metadata."""
        wb = self.create_test_workbook()

        metadata = self.extractor.extract_key_value_pairs(wb, 'TestSheet', 'A1:B5')

        self.assertIn('Equipment ID', metadata)
        self.assertEqual(metadata['Equipment ID'], 'TEST-123')
        self.assertIn('Test Date', metadata)

    def test_extract_merged_cells(self):
        """Test extracting merged cell ranges."""
        wb = Workbook()
        ws = wb.active

        ws.merge_cells('A1:B1')
        ws['A1'] = 'Merged Header'

        merged = self.extractor.extract_merged_cells(wb, ws.title)

        self.assertEqual(len(merged), 1)
        self.assertIn('A1:B1', merged[0])


class TestExcelTransformer(unittest.TestCase):
    """Test cases for ExcelTransformer class."""

    def setUp(self):
        """Set up test fixtures."""
        self.transformer = ExcelTransformer()

    def test_convert_units_voltage(self):
        """Test voltage unit conversion."""
        df = pd.DataFrame({
            'Voltage (mV)': [37500, 38000, 37800]
        })

        df = self.transformer.convert_units(
            df, 'Voltage (mV)', 'mV', 'V', 'voltage'
        )

        self.assertIn('Voltage (V)', df.columns)
        self.assertAlmostEqual(df['Voltage (V)'].iloc[0], 37.5, places=1)

    def test_convert_units_current(self):
        """Test current unit conversion."""
        df = pd.DataFrame({
            'Current (mA)': [8000, 8100, 7900]
        })

        df = self.transformer.convert_units(
            df, 'Current (mA)', 'mA', 'A', 'current'
        )

        self.assertIn('Current (A)', df.columns)
        self.assertAlmostEqual(df['Current (A)'].iloc[0], 8.0, places=1)

    def test_clean_column_names(self):
        """Test column name cleaning."""
        df = pd.DataFrame({
            'Column With Spaces': [1, 2, 3],
            'Column-With-Dashes': [4, 5, 6],
            'Column (with parens)': [7, 8, 9]
        })

        df = self.transformer.clean_column_names(
            df,
            lowercase=True,
            remove_special_chars=True,
            replace_spaces='_'
        )

        self.assertIn('column_with_spaces', df.columns)

    def test_parse_timestamp_excel_serial(self):
        """Test Excel serial date parsing."""
        # Excel serial date for 2024-01-01 is approximately 45292
        df = pd.DataFrame({
            'date': [45292, 45293, 45294]
        })

        df = self.transformer.parse_timestamp(df, 'date', handle_excel_serial=True)

        self.assertTrue(pd.api.types.is_datetime64_any_dtype(df['date']))

    def test_fill_missing_values_forward(self):
        """Test forward fill of missing values."""
        df = pd.DataFrame({
            'A': [1, None, None, 4, 5]
        })

        df = self.transformer.fill_missing_values(df, method='forward')

        self.assertEqual(df['A'].iloc[1], 1.0)
        self.assertEqual(df['A'].iloc[2], 1.0)

    def test_extract_unit_from_column_name(self):
        """Test unit extraction from column names."""
        test_cases = [
            ('Temperature (°C)', ('Temperature', '°C')),
            ('Voltage [V]', ('Voltage', 'V')),
            ('Current_A', ('Current', 'A')),
            ('Plain Column', ('Plain Column', None))
        ]

        for col_name, expected in test_cases:
            base, unit = self.transformer.extract_unit_from_column_name(col_name)
            self.assertEqual((base, unit), expected)

    def test_split_datetime_components(self):
        """Test datetime component splitting."""
        df = pd.DataFrame({
            'timestamp': pd.date_range('2024-01-01', periods=5, freq='D')
        })

        df = self.transformer.split_datetime_components(
            df,
            'timestamp',
            components=['year', 'month', 'day']
        )

        self.assertIn('timestamp_year', df.columns)
        self.assertIn('timestamp_month', df.columns)
        self.assertIn('timestamp_day', df.columns)
        self.assertEqual(df['timestamp_year'].iloc[0], 2024)
        self.assertEqual(df['timestamp_month'].iloc[0], 1)


class TestExcelIngestion(unittest.TestCase):
    """Test cases for ExcelIngestion class."""

    def setUp(self):
        """Set up test fixtures."""
        self.temp_dir = tempfile.mkdtemp()
        self.test_file = Path(self.temp_dir) / 'test.xlsx'

    def tearDown(self):
        """Clean up temporary files."""
        import shutil
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def create_simple_test_file(self):
        """Create a simple test Excel file."""
        wb = Workbook()
        ws = wb.active
        ws.title = "TestData"

        ws['A1'] = 'Column1'
        ws['B1'] = 'Column2'
        ws['A2'] = 1
        ws['B2'] = 'a'
        ws['A3'] = 2
        ws['B3'] = 'b'

        wb.save(self.test_file)

    def test_file_not_found(self):
        """Test exception when file doesn't exist."""
        with self.assertRaises(FileNotFoundError):
            ExcelIngestion('nonexistent.xlsx')

    def test_file_hash_calculation(self):
        """Test SHA-256 file hash calculation."""
        self.create_simple_test_file()

        ingestion = ExcelIngestion(self.test_file)

        self.assertIsNotNone(ingestion.file_hash)
        self.assertEqual(len(ingestion.file_hash), 64)  # SHA-256 hex length

    def test_get_sheet_names(self):
        """Test retrieving sheet names."""
        self.create_simple_test_file()

        ingestion = ExcelIngestion(self.test_file)
        sheets = ingestion.get_sheet_names()

        self.assertIn('TestData', sheets)

    def test_extract_metadata(self):
        """Test metadata extraction."""
        self.create_simple_test_file()

        ingestion = ExcelIngestion(self.test_file)
        metadata = ingestion.extract_metadata()

        self.assertIn('file_name', metadata)
        self.assertIn('file_hash', metadata)
        self.assertIn('sheet_count', metadata)
        self.assertEqual(metadata['file_name'], 'test.xlsx')

    def test_parse_sheet(self):
        """Test basic sheet parsing."""
        self.create_simple_test_file()

        ingestion = ExcelIngestion(self.test_file)
        result = ingestion.parse_sheet('TestData')

        self.assertTrue(result.is_valid)
        self.assertIsNotNone(result.dataframe)
        self.assertEqual(len(result.dataframe), 2)

    def test_parse_invalid_sheet(self):
        """Test parsing non-existent sheet."""
        self.create_simple_test_file()

        ingestion = ExcelIngestion(self.test_file)

        with self.assertRaises(InvalidSheetError):
            ingestion.parse_sheet('NonExistentSheet')

    def test_parse_iec61215_mst(self):
        """Test IEC 61215 MST data parsing."""
        # Create MST test file
        wb = Workbook()
        ws = wb.active
        ws.title = "DampHeat"

        # Headers
        ws['A10'] = 'Time (h)'
        ws['B10'] = 'Temperature (°C)'
        ws['C10'] = 'Humidity (%RH)'

        # Data (IEC 61215: 85°C, 85% RH)
        for i in range(10):
            ws[f'A{11+i}'] = i
            ws[f'B{11+i}'] = 85 + np.random.uniform(-0.5, 0.5)
            ws[f'C{11+i}'] = 85 + np.random.uniform(-1, 1)

        wb.save(self.test_file)

        ingestion = ExcelIngestion(self.test_file)
        result = ingestion.parse_iec61215_mst(
            sheet_name='DampHeat',
            data_range='A10:C20'
        )

        self.assertTrue(result.is_valid)
        self.assertEqual(result.metadata['test_type'], 'IEC 61215 MST Damp Heat')
        self.assertIn('avg_temperature', result.metadata)
        self.assertIn('avg_humidity', result.metadata)

    def test_parse_iv_curve(self):
        """Test I-V curve data parsing."""
        wb = Workbook()
        ws = wb.active
        ws.title = "IV_Curve"

        # Metadata
        ws['A1'] = 'Module Serial:'
        ws['B1'] = 'TEST-MODULE'

        # Headers
        ws['A5'] = 'Voltage (V)'
        ws['B5'] = 'Current (A)'

        # Data
        voltages = np.linspace(0, 45, 20)
        currents = 9 * (1 - voltages / 45)

        for i, (v, i_val) in enumerate(zip(voltages, currents)):
            ws[f'A{6+i}'] = v
            ws[f'B{6+i}'] = max(0, i_val)

        wb.save(self.test_file)

        ingestion = ExcelIngestion(self.test_file)
        result = ingestion.parse_iv_curve(
            sheet_name='IV_Curve',
            data_range='A5:B25'
        )

        self.assertTrue(result.is_valid)
        self.assertIn('voc', result.metadata)
        self.assertIn('isc', result.metadata)
        self.assertIn('pmax', result.metadata)
        self.assertIn('fill_factor', result.metadata)


class TestIntegration(unittest.TestCase):
    """Integration tests for end-to-end workflows."""

    def setUp(self):
        """Set up test fixtures."""
        self.temp_dir = tempfile.mkdtemp()

    def tearDown(self):
        """Clean up temporary files."""
        import shutil
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_full_workflow_mst_data(self):
        """Test complete workflow: load, parse, validate, transform."""
        # Create test file
        test_file = Path(self.temp_dir) / 'mst_test.xlsx'

        wb = Workbook()
        ws = wb.active
        ws.title = "DampHeat"

        # Add data
        ws['A1'] = 'Time (h)'
        ws['B1'] = 'Temperature (°C)'
        ws['C1'] = 'Humidity (%RH)'

        np.random.seed(42)
        for i in range(100):
            ws[f'A{2+i}'] = i
            ws[f'B{2+i}'] = 85 + np.random.normal(0, 0.5)
            ws[f'C{2+i}'] = 85 + np.random.normal(0, 1.0)

        wb.save(test_file)

        # Parse
        ingestion = ExcelIngestion(test_file)
        result = ingestion.parse_iec61215_mst(sheet_name='DampHeat')

        # Validate result
        self.assertTrue(result.is_valid)
        self.assertIsNotNone(result.dataframe)
        self.assertEqual(len(result.dataframe), 100)

        # Transform
        transformer = ExcelTransformer()
        df = result.dataframe

        # Clean column names
        df = transformer.clean_column_names(df)

        # Check transformations
        self.assertTrue(all('_' in col or col.isalnum() for col in df.columns))

    def test_unit_conversion_workflow(self):
        """Test workflow with unit conversion."""
        test_file = Path(self.temp_dir) / 'voltage_test.xlsx'

        wb = Workbook()
        ws = wb.active

        # Data in mV (should be converted to V)
        ws['A1'] = 'Voltage (mV)'
        for i in range(10):
            ws[f'A{2+i}'] = 37000 + i * 100

        wb.save(test_file)

        # Parse
        ingestion = ExcelIngestion(test_file)
        result = ingestion.parse_sheet(ws.title)

        # Transform
        transformer = ExcelTransformer()
        df = result.dataframe

        df = transformer.convert_units(df, 'Voltage (mV)', 'mV', 'V', 'voltage')

        # Validate conversion
        self.assertIn('Voltage (V)', df.columns)
        self.assertAlmostEqual(df['Voltage (V)'].iloc[0], 37.0, places=1)


def run_tests():
    """Run all unit tests."""
    # Create test suite
    loader = unittest.TestLoader()
    suite = unittest.TestSuite()

    # Add all test cases
    suite.addTests(loader.loadTestsFromTestCase(TestExcelValidator))
    suite.addTests(loader.loadTestsFromTestCase(TestExcelExtractor))
    suite.addTests(loader.loadTestsFromTestCase(TestExcelTransformer))
    suite.addTests(loader.loadTestsFromTestCase(TestExcelIngestion))
    suite.addTests(loader.loadTestsFromTestCase(TestIntegration))

    # Run tests
    runner = unittest.TextTestRunner(verbosity=2)
    result = runner.run(suite)

    return result.wasSuccessful()


if __name__ == '__main__':
    success = run_tests()
    exit(0 if success else 1)
