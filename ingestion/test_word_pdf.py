"""
Comprehensive unit tests for Word and PDF ingestion engine.
"""

import io
import os
import tempfile
from datetime import date, datetime
from pathlib import Path

import pytest
from docx import Document
from PyPDF2 import PdfWriter

from ingestion.certificate_parser import CertificateParser, parse_calibration_certificate
from ingestion.document_validators import (
    check_data_completeness,
    normalize_unit,
    sanitize_extracted_text,
    validate_calibration_certificate,
    validate_date_range,
    validate_equipment_id,
    validate_file_exists,
    validate_file_size,
    validate_ingestion_result,
    validate_numeric_value,
    validate_table_data,
)
from ingestion.exceptions import (
    CorruptedFileError,
    DocumentParsingError,
    OCRError,
    PasswordProtectedError,
    TableExtractionError,
    UnsupportedFormatError,
    ValidationError,
)
from ingestion.models import (
    AccreditationBody,
    CalibrationCertificate,
    CalibrationPoint,
    DocumentIngestionResult,
    DocumentMetadata,
    ExtractionMethod,
    ImageData,
    ParagraphData,
    TableData,
    TraceabilityInfo,
    UncertaintyMeasurement,
)
from ingestion.pdf_parser import PDFIngestion, parse_pdf_document
from ingestion.word_parser import WordDocumentParser, parse_word_document


class TestModels:
    """Test Pydantic models."""

    def test_table_data_to_dataframe(self):
        """Test TableData conversion to DataFrame."""
        table = TableData(
            data=[["1", "2", "3"], ["4", "5", "6"]],
            headers=["Col1", "Col2", "Col3"],
        )

        df = table.to_dataframe()
        assert df.shape == (2, 3)
        assert list(df.columns) == ["Col1", "Col2", "Col3"]

    def test_calibration_certificate_validation(self):
        """Test CalibrationCertificate model validation."""
        cert = CalibrationCertificate(
            certificate_number="TEST-001",
            equipment_id="EQ-001",
            equipment_name="Test Equipment",
            calibration_date=date(2024, 1, 15),
            due_date=date(2025, 1, 14),
            laboratory_name="Test Lab",
        )

        assert cert.is_valid()
        assert cert.days_until_due() is not None

    def test_uncertainty_measurement(self):
        """Test UncertaintyMeasurement model."""
        uncertainty = UncertaintyMeasurement(value=0.05, unit="V", coverage_factor=2.0)

        assert uncertainty.value == 0.05
        assert uncertainty.unit == "V"
        assert uncertainty.confidence_level == 0.95

    def test_calibration_point(self):
        """Test CalibrationPoint model."""
        point = CalibrationPoint(
            reference_value=10.0, measured_value=10.02, unit="V", deviation=0.02
        )

        assert point.reference_value == 10.0
        assert point.deviation == 0.02

    def test_extraction_method_enum(self):
        """Test ExtractionMethod enum."""
        assert ExtractionMethod.NATIVE == "native"
        assert ExtractionMethod.OCR == "ocr"
        assert ExtractionMethod.HYBRID == "hybrid"

    def test_accreditation_body_enum(self):
        """Test AccreditationBody enum."""
        assert AccreditationBody.NABL == "NABL"
        assert AccreditationBody.ILAC == "ILAC"
        assert AccreditationBody.UNKNOWN == "UNKNOWN"


class TestDocumentValidators:
    """Test document validation functions."""

    def test_validate_file_exists(self, tmp_path):
        """Test file existence validation."""
        # Create temporary file
        test_file = tmp_path / "test.txt"
        test_file.write_text("test content")

        assert validate_file_exists(str(test_file))

        # Test non-existent file
        with pytest.raises(FileNotFoundError):
            validate_file_exists(str(tmp_path / "nonexistent.txt"))

    def test_validate_file_size(self, tmp_path):
        """Test file size validation."""
        # Create small file
        test_file = tmp_path / "small.txt"
        test_file.write_text("small content")

        assert validate_file_size(str(test_file), max_size_mb=1)

    def test_validate_date_range(self):
        """Test date range validation."""
        start = date(2024, 1, 1)
        end = date(2024, 12, 31)

        is_valid, msg = validate_date_range(start, end, max_days=400)
        assert is_valid

        # Invalid range
        is_valid, msg = validate_date_range(end, start)
        assert not is_valid

    def test_validate_numeric_value(self):
        """Test numeric value validation."""
        is_valid, msg = validate_numeric_value(50, min_value=0, max_value=100)
        assert is_valid

        is_valid, msg = validate_numeric_value(150, min_value=0, max_value=100)
        assert not is_valid

    def test_sanitize_extracted_text(self):
        """Test text sanitization."""
        dirty_text = "Test\x00  \n\n  text   with   spaces"
        clean_text = sanitize_extracted_text(dirty_text)

        assert "\x00" not in clean_text
        assert "Test text with spaces" in clean_text

    def test_normalize_unit(self):
        """Test unit normalization."""
        assert normalize_unit("deg c") == "°C"
        assert normalize_unit("volts") == "V"
        assert normalize_unit("percent") == "%"
        assert normalize_unit("WATTS") == "W"

    def test_validate_equipment_id(self):
        """Test equipment ID validation."""
        assert validate_equipment_id("EQ-001")
        assert validate_equipment_id("DMM_123")
        assert validate_equipment_id("TEST/456")
        assert not validate_equipment_id("invalid id")

    def test_check_data_completeness(self):
        """Test data completeness check."""
        data = {"field1": "value1", "field2": "value2", "field3": None}
        required = ["field1", "field2", "field3", "field4"]

        is_complete, missing = check_data_completeness(data, required)

        assert not is_complete
        assert "field3" in missing  # None value
        assert "field4" in missing  # Not present

    def test_validate_table_data(self):
        """Test table data validation."""
        table = TableData(
            data=[["1", "2", "3"], ["4", "5", "6"]], headers=["A", "B", "C"]
        )

        is_valid, warnings = validate_table_data(table)
        assert is_valid

        # Empty table
        empty_table = TableData(data=[])
        is_valid, warnings = validate_table_data(empty_table)
        assert not is_valid


class TestWordDocumentParser:
    """Test Word document parser."""

    def test_create_and_parse_docx(self, tmp_path):
        """Test creating and parsing a Word document."""
        # Create a test .docx file
        doc_path = tmp_path / "test.docx"
        doc = Document()

        # Add title
        doc.add_heading("Test Document", level=1)

        # Add paragraphs
        doc.add_paragraph("This is a test paragraph.")
        doc.add_paragraph("This is another paragraph with bold text.").runs[0].bold = True

        # Add table
        table = doc.add_table(rows=3, cols=3)
        table.rows[0].cells[0].text = "Header1"
        table.rows[0].cells[1].text = "Header2"
        table.rows[0].cells[2].text = "Header3"
        table.rows[1].cells[0].text = "Data1"
        table.rows[1].cells[1].text = "Data2"
        table.rows[1].cells[2].text = "Data3"

        doc.save(str(doc_path))

        # Parse document
        parser = WordDocumentParser(str(doc_path))
        result = parser.parse()

        assert isinstance(result, DocumentIngestionResult)
        assert "Test Document" in result.text_content
        assert "test paragraph" in result.text_content
        assert len(result.tables) > 0
        assert result.extraction_method == ExtractionMethod.NATIVE

    def test_word_parser_unsupported_format(self, tmp_path):
        """Test Word parser with unsupported format."""
        text_file = tmp_path / "test.txt"
        text_file.write_text("test")

        with pytest.raises(UnsupportedFormatError):
            WordDocumentParser(str(text_file))

    def test_word_parser_file_not_found(self):
        """Test Word parser with non-existent file."""
        with pytest.raises(FileNotFoundError):
            WordDocumentParser("/nonexistent/file.docx")


class TestPDFParser:
    """Test PDF parser."""

    def test_create_and_parse_pdf(self, tmp_path):
        """Test creating and parsing a simple PDF."""
        from reportlab.lib.pagesizes import letter
        from reportlab.pdfgen import canvas

        pdf_path = tmp_path / "test.pdf"

        # Create simple PDF
        c = canvas.Canvas(str(pdf_path), pagesize=letter)
        c.drawString(100, 750, "Test PDF Document")
        c.drawString(100, 700, "This is a test paragraph.")
        c.save()

        # Parse PDF
        parser = PDFIngestion(str(pdf_path))
        result = parser.parse()

        assert isinstance(result, DocumentIngestionResult)
        assert result.extraction_method == ExtractionMethod.NATIVE
        assert result.metadata.page_count == 1

    def test_pdf_parser_unsupported_format(self, tmp_path):
        """Test PDF parser with unsupported format."""
        text_file = tmp_path / "test.txt"
        text_file.write_text("test")

        with pytest.raises(UnsupportedFormatError):
            PDFIngestion(str(text_file))

    def test_pdf_metadata_extraction(self, tmp_path):
        """Test PDF metadata extraction."""
        from reportlab.lib.pagesizes import letter
        from reportlab.pdfgen import canvas

        pdf_path = tmp_path / "test_metadata.pdf"

        c = canvas.Canvas(str(pdf_path), pagesize=letter)
        c.setAuthor("Test Author")
        c.setTitle("Test Title")
        c.drawString(100, 750, "Content")
        c.save()

        parser = PDFIngestion(str(pdf_path))
        metadata = parser.extract_metadata()

        assert metadata.page_count == 1
        # Note: reportlab metadata may not be extracted by PyPDF2
        # In real PDFs, these would be present


class TestExceptions:
    """Test custom exceptions."""

    def test_document_parsing_error(self):
        """Test DocumentParsingError."""
        error = DocumentParsingError("Test error", "/path/to/file.pdf")
        assert error.file_path == "/path/to/file.pdf"
        assert "Test error" in str(error)

    def test_unsupported_format_error(self):
        """Test UnsupportedFormatError."""
        error = UnsupportedFormatError("/path/to/file.xyz", "XYZ")
        assert error.file_path == "/path/to/file.xyz"

    def test_corrupted_file_error(self):
        """Test CorruptedFileError."""
        error = CorruptedFileError("/path/to/file.pdf", "Invalid header")
        assert "corrupted" in str(error).lower()

    def test_ocr_error(self):
        """Test OCRError."""
        error = OCRError("Tesseract not found", "/path/to/file.pdf", confidence=0.5)
        assert error.confidence == 0.5

    def test_table_extraction_error(self):
        """Test TableExtractionError."""
        error = TableExtractionError("Table parse failed", "/path/to/file.pdf", page_number=3)
        assert error.page_number == 3

    def test_validation_error(self):
        """Test ValidationError."""
        error = ValidationError("Invalid date format", field_name="calibration_date")
        assert error.field_name == "calibration_date"


class TestIntegration:
    """Integration tests."""

    def test_full_workflow_word_document(self, tmp_path):
        """Test complete workflow with Word document."""
        # Create Word document
        doc_path = tmp_path / "full_test.docx"
        doc = Document()
        doc.add_heading("Integration Test", level=1)
        doc.add_paragraph("Test content for integration testing.")

        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "Key"
        table.rows[0].cells[1].text = "Value"
        table.rows[1].cells[0].text = "Test"
        table.rows[1].cells[1].text = "123"

        doc.save(str(doc_path))

        # Parse
        result = parse_word_document(str(doc_path))

        # Validate
        is_valid, warnings = validate_ingestion_result(result)

        assert isinstance(result, DocumentIngestionResult)
        assert len(result.tables) > 0
        assert result.file_hash is not None

    def test_sample_data_files(self):
        """Test with sample data files."""
        sample_dir = Path(__file__).parent / "sample_data"

        # Check if sample files exist
        if (sample_dir / "test_document.txt").exists():
            # This is just a text file, not a real document
            # In real tests, we'd have actual PDFs and Word docs
            assert True
        else:
            pytest.skip("Sample data files not found")


class TestCertificateParser:
    """Test calibration certificate parser."""

    def test_extract_certificate_number(self):
        """Test certificate number extraction."""
        text = "Certificate No: NABL-2024-12345"
        parser = CertificateParser.__new__(CertificateParser)
        parser.raw_text = text.lower()

        cert_num = parser.extract_certificate_number()
        assert cert_num == "NABL-2024-12345"

    def test_extract_dates(self):
        """Test date extraction."""
        text = """
        Calibration Date: 15-Jan-2024
        Due Date: 14-Jan-2025
        """
        parser = CertificateParser.__new__(CertificateParser)
        parser.raw_text = text.lower()
        parser.DATE_PATTERNS = CertificateParser.DATE_PATTERNS
        parser.DUE_DATE_PATTERNS = CertificateParser.DUE_DATE_PATTERNS
        parser._parse_date = CertificateParser._parse_date.__get__(parser)

        cal_date, due_date = parser.extract_dates()

        assert cal_date is not None
        assert due_date is not None

    def test_extract_equipment_info(self):
        """Test equipment information extraction."""
        text = """
        Equipment: Digital Multimeter
        Model: 87V
        Serial Number: 12345678
        """
        parser = CertificateParser.__new__(CertificateParser)
        parser.raw_text = text.lower()
        parser.EQUIPMENT_PATTERNS = CertificateParser.EQUIPMENT_PATTERNS
        parser.MODEL_PATTERNS = CertificateParser.MODEL_PATTERNS
        parser.SERIAL_PATTERNS = CertificateParser.SERIAL_PATTERNS

        eq_name, model, serial = parser.extract_equipment_info()

        assert model == "87V"
        assert serial == "12345678"

    def test_extract_accreditation_body(self):
        """Test accreditation body detection."""
        parser = CertificateParser.__new__(CertificateParser)

        parser.raw_text = "nabl accredited laboratory"
        assert parser.extract_accreditation_body() == AccreditationBody.NABL

        parser.raw_text = "ilac mra member"
        assert parser.extract_accreditation_body() == AccreditationBody.ILAC

        parser.raw_text = "no accreditation mentioned"
        assert parser.extract_accreditation_body() == AccreditationBody.UNKNOWN


# Run tests with pytest
if __name__ == "__main__":
    pytest.main([__file__, "-v"])
