"""
File parser module for PV Test Report Automation
Handles parsing of Excel, PDF, images, and Word documents
"""

import pandas as pd
from PIL import Image
from datetime import datetime
from typing import Dict, List, Optional, Any
import io
from pathlib import Path


class FileParser:
    """Base class for file parsing operations"""

    SUPPORTED_EXTENSIONS = {
        'excel': ['.xlsx', '.xls', '.xlsm'],
        'image': ['.jpg', '.jpeg', '.png', '.gif', '.bmp', '.tiff'],
        'pdf': ['.pdf'],
        'word': ['.docx', '.doc']
    }

    @staticmethod
    def get_file_type(filename: str) -> Optional[str]:
        """Determine file type based on extension"""
        ext = Path(filename).suffix.lower()
        for file_type, extensions in FileParser.SUPPORTED_EXTENSIONS.items():
            if ext in extensions:
                return file_type
        return None

    @staticmethod
    def validate_file(file_obj, filename: str) -> Dict[str, Any]:
        """
        Validate uploaded file

        Args:
            file_obj: File object from Streamlit uploader
            filename: Name of the file

        Returns:
            Dict with validation results
        """
        result = {
            'valid': False,
            'file_type': None,
            'size': 0,
            'error': None
        }

        try:
            file_type = FileParser.get_file_type(filename)
            if not file_type:
                result['error'] = f"Unsupported file type: {Path(filename).suffix}"
                return result

            # Get file size
            file_obj.seek(0, 2)  # Seek to end
            size = file_obj.tell()
            file_obj.seek(0)  # Reset to beginning

            result['valid'] = True
            result['file_type'] = file_type
            result['size'] = size

        except Exception as e:
            result['error'] = str(e)

        return result


class ExcelParser(FileParser):
    """Parser for Excel files"""

    @staticmethod
    def parse(file_obj, filename: str) -> Dict[str, Any]:
        """
        Parse Excel file and extract data

        Args:
            file_obj: File object from Streamlit uploader
            filename: Name of the file

        Returns:
            Dict containing parsed data and metadata
        """
        result = {
            'success': False,
            'sheets': {},
            'sheet_names': [],
            'metadata': {},
            'error': None
        }

        try:
            # Read Excel file
            excel_file = pd.ExcelFile(file_obj)
            result['sheet_names'] = excel_file.sheet_names

            # Parse each sheet
            for sheet_name in excel_file.sheet_names:
                df = pd.read_excel(excel_file, sheet_name=sheet_name)
                result['sheets'][sheet_name] = {
                    'data': df,
                    'rows': len(df),
                    'columns': len(df.columns),
                    'column_names': df.columns.tolist()
                }

            # Metadata
            result['metadata'] = {
                'filename': filename,
                'total_sheets': len(excel_file.sheet_names),
                'parsed_at': datetime.now().isoformat()
            }

            result['success'] = True

        except Exception as e:
            result['error'] = str(e)

        return result

    @staticmethod
    def get_sheet_preview(df: pd.DataFrame, num_rows: int = 5) -> pd.DataFrame:
        """Get preview of dataframe"""
        return df.head(num_rows)


class ImageParser(FileParser):
    """Parser for image files"""

    @staticmethod
    def parse(file_obj, filename: str) -> Dict[str, Any]:
        """
        Parse image file and extract metadata

        Args:
            file_obj: File object from Streamlit uploader
            filename: Name of the file

        Returns:
            Dict containing image metadata
        """
        result = {
            'success': False,
            'metadata': {},
            'error': None
        }

        try:
            # Open image
            image = Image.open(file_obj)

            # Extract metadata
            result['metadata'] = {
                'filename': filename,
                'format': image.format,
                'mode': image.mode,
                'size': image.size,
                'width': image.width,
                'height': image.height,
                'file_size': file_obj.tell(),
                'parsed_at': datetime.now().isoformat()
            }

            # Get EXIF data if available
            exif_data = image.getexif()
            if exif_data:
                result['metadata']['exif'] = dict(exif_data)

            result['success'] = True

        except Exception as e:
            result['error'] = str(e)

        return result


class PDFParser(FileParser):
    """Parser for PDF files"""

    @staticmethod
    def parse(file_obj, filename: str) -> Dict[str, Any]:
        """
        Parse PDF file and extract basic metadata

        Args:
            file_obj: File object from Streamlit uploader
            filename: Name of the file

        Returns:
            Dict containing PDF metadata
        """
        result = {
            'success': False,
            'metadata': {},
            'error': None
        }

        try:
            import PyPDF2

            # Read PDF
            pdf_reader = PyPDF2.PdfReader(file_obj)

            # Extract metadata
            result['metadata'] = {
                'filename': filename,
                'num_pages': len(pdf_reader.pages),
                'parsed_at': datetime.now().isoformat()
            }

            # Get PDF info if available
            if pdf_reader.metadata:
                result['metadata']['info'] = {
                    key: str(value) for key, value in pdf_reader.metadata.items()
                }

            result['success'] = True

        except Exception as e:
            result['error'] = str(e)

        return result


class WordParser(FileParser):
    """Parser for Word documents"""

    @staticmethod
    def parse(file_obj, filename: str) -> Dict[str, Any]:
        """
        Parse Word document and extract basic metadata

        Args:
            file_obj: File object from Streamlit uploader
            filename: Name of the file

        Returns:
            Dict containing document metadata
        """
        result = {
            'success': False,
            'metadata': {},
            'error': None
        }

        try:
            from docx import Document

            # Read document
            doc = Document(file_obj)

            # Extract metadata
            result['metadata'] = {
                'filename': filename,
                'num_paragraphs': len(doc.paragraphs),
                'num_tables': len(doc.tables),
                'parsed_at': datetime.now().isoformat()
            }

            # Get core properties if available
            if doc.core_properties:
                props = doc.core_properties
                result['metadata']['properties'] = {
                    'author': props.author,
                    'title': props.title,
                    'subject': props.subject,
                    'created': str(props.created) if props.created else None,
                    'modified': str(props.modified) if props.modified else None
                }

            result['success'] = True

        except Exception as e:
            result['error'] = str(e)

        return result


def parse_file(file_obj, filename: str) -> Dict[str, Any]:
    """
    Main function to parse any supported file type

    Args:
        file_obj: File object from Streamlit uploader
        filename: Name of the file

    Returns:
        Dict containing parsed data and metadata
    """
    # Validate file first
    validation = FileParser.validate_file(file_obj, filename)

    if not validation['valid']:
        return {
            'success': False,
            'error': validation['error'],
            'file_type': None
        }

    file_type = validation['file_type']

    # Parse based on file type
    if file_type == 'excel':
        return ExcelParser.parse(file_obj, filename)
    elif file_type == 'image':
        return ImageParser.parse(file_obj, filename)
    elif file_type == 'pdf':
        return PDFParser.parse(file_obj, filename)
    elif file_type == 'word':
        return WordParser.parse(file_obj, filename)
    else:
        return {
            'success': False,
            'error': f'No parser available for file type: {file_type}',
            'file_type': file_type
        }
