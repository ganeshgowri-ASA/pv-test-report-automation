"""
Word Exporter Module

Provides Word document export functionality using python-docx.
Includes template support, tables, images, and styling.
"""

from pathlib import Path
from typing import Dict, Any, Optional, List, TYPE_CHECKING
from datetime import datetime
import logging
import io
import base64

try:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    if TYPE_CHECKING:
        from typing import Any as RGBColor
        Document = Any
    else:
        RGBColor = Any
        Document = object

from .base_exporter import (
    BaseExporter, ExportFormat, ExportOptions, ExportRegistry
)


logger = logging.getLogger(__name__)


@ExportRegistry.register(ExportFormat.WORD)
class WordExporter(BaseExporter):
    """Word document export engine."""

    def __init__(self, template_manager=None):
        """
        Initialize Word exporter.

        Args:
            template_manager: Template manager instance
        """
        super().__init__(template_manager)
        self._validate_dependencies()

    @property
    def format_type(self) -> ExportFormat:
        """Get export format type."""
        return ExportFormat.WORD

    @property
    def file_extension(self) -> str:
        """Get file extension."""
        return "docx"

    def _validate_dependencies(self):
        """Validate required dependencies."""
        if not DOCX_AVAILABLE:
            raise ImportError("python-docx not installed. Install with: pip install python-docx")

    def export(self, data: Dict[str, Any], options: ExportOptions) -> Path:
        """
        Export data to Word document.

        Args:
            data: Data to export
            options: Export options

        Returns:
            Path to exported Word file
        """
        self._validate_data(data)
        self._ensure_output_directory(options.output_path)

        logger.info(f"Exporting Word document: {options.output_path}")

        # Add metadata
        if options.include_metadata:
            data = self._add_metadata(data)

        # Load template or create new document
        if options.template_name and self.template_manager:
            template_path = self.template_manager.get_template(
                options.template_name, self.format_type
            )
            if template_path.exists():
                doc = Document(str(template_path))
            else:
                doc = Document()
        else:
            doc = Document()

        # Setup document styles
        self._setup_styles(doc)

        # Set document properties
        self._set_document_properties(doc, data)

        # Build document content
        self._build_document(doc, data, options)

        # Save document
        doc.save(str(options.output_path))

        # Validate output
        if not self.validate_output(options.output_path):
            raise ValueError(f"Invalid Word output: {options.output_path}")

        logger.info(f"Word export completed: {options.output_path}")
        return options.output_path

    def _setup_styles(self, doc: 'Document'):
        """
        Setup custom document styles.

        Args:
            doc: Document instance
        """
        styles = doc.styles

        # Custom title style
        if 'CustomTitle' not in styles:
            title_style = styles.add_style('CustomTitle', WD_STYLE_TYPE.PARAGRAPH)
            title_font = title_style.font
            title_font.name = 'Calibri'
            title_font.size = Pt(28)
            title_font.bold = True
            title_font.color.rgb = RGBColor(31, 71, 136)  # #1f4788
            title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            title_style.paragraph_format.space_after = Pt(20)

        # Custom heading style
        if 'CustomHeading' not in styles:
            heading_style = styles.add_style('CustomHeading', WD_STYLE_TYPE.PARAGRAPH)
            heading_font = heading_style.font
            heading_font.name = 'Calibri'
            heading_font.size = Pt(16)
            heading_font.bold = True
            heading_font.color.rgb = RGBColor(45, 90, 160)  # #2d5aa0
            heading_style.paragraph_format.space_before = Pt(12)
            heading_style.paragraph_format.space_after = Pt(6)

        # Custom body style
        if 'CustomBody' not in styles:
            body_style = styles.add_style('CustomBody', WD_STYLE_TYPE.PARAGRAPH)
            body_font = body_style.font
            body_font.name = 'Calibri'
            body_font.size = Pt(11)

    def _set_document_properties(self, doc: 'Document', data: Dict[str, Any]):
        """
        Set document properties.

        Args:
            doc: Document instance
            data: Report data
        """
        core_properties = doc.core_properties
        core_properties.title = data.get("title", "PV Test Report")
        core_properties.author = data.get("author", "PV Test Automation System")
        core_properties.subject = "PV Test Report"
        core_properties.comments = data.get("description", "Automated test report")
        core_properties.created = datetime.now()

    def _build_document(self, doc: 'Document', data: Dict[str, Any], options: ExportOptions):
        """
        Build document content.

        Args:
            doc: Document instance
            data: Report data
            options: Export options
        """
        # Title
        title = data.get("title", "PV Test Report")
        doc.add_paragraph(title, style='CustomTitle')

        # Metadata section
        if options.include_metadata and "_export_metadata" in data:
            self._add_metadata_section(doc, data["_export_metadata"])

        # Executive Summary
        if "summary" in data:
            self._add_summary_section(doc, data["summary"])

        # Test Configuration
        if "configuration" in data:
            self._add_configuration_section(doc, data["configuration"])

        # Test Results
        if "test_results" in data:
            self._add_test_results_section(doc, data["test_results"])

        # Charts and Graphs
        if options.include_charts and "charts" in data:
            self._add_charts_section(doc, data["charts"])

        # Analysis
        if "analysis" in data:
            self._add_analysis_section(doc, data["analysis"])

        # Recommendations
        if "recommendations" in data:
            self._add_recommendations_section(doc, data["recommendations"])

        # Appendix
        if "appendix" in data:
            self._add_appendix_section(doc, data["appendix"])

    def _add_metadata_section(self, doc: 'Document', metadata: Dict[str, Any]):
        """
        Add metadata section.

        Args:
            doc: Document instance
            metadata: Metadata dictionary
        """
        doc.add_paragraph("Document Information", style='CustomHeading')

        table = doc.add_table(rows=len(metadata), cols=2)
        table.style = 'Light Grid Accent 1'

        for idx, (key, value) in enumerate(metadata.items()):
            row = table.rows[idx]
            row.cells[0].text = key.replace("_", " ").title()
            row.cells[1].text = str(value)

            # Bold the first column
            row.cells[0].paragraphs[0].runs[0].font.bold = True

        doc.add_paragraph()  # Add spacing

    def _add_summary_section(self, doc: 'Document', summary: str):
        """
        Add executive summary section.

        Args:
            doc: Document instance
            summary: Summary text
        """
        doc.add_paragraph("Executive Summary", style='CustomHeading')
        para = doc.add_paragraph(summary, style='CustomBody')
        para.paragraph_format.space_after = Pt(12)
        doc.add_paragraph()

    def _add_configuration_section(self, doc: 'Document', configuration: Dict[str, Any]):
        """
        Add test configuration section.

        Args:
            doc: Document instance
            configuration: Configuration data
        """
        doc.add_paragraph("Test Configuration", style='CustomHeading')

        table = doc.add_table(rows=len(configuration), cols=2)
        table.style = 'Light Grid Accent 1'

        for idx, (key, value) in enumerate(configuration.items()):
            row = table.rows[idx]
            row.cells[0].text = key.replace("_", " ").title()
            row.cells[1].text = str(value)
            row.cells[0].paragraphs[0].runs[0].font.bold = True

        doc.add_paragraph()

    def _add_test_results_section(self, doc: 'Document', test_results: List[Dict[str, Any]]):
        """
        Add test results section.

        Args:
            doc: Document instance
            test_results: List of test results
        """
        doc.add_paragraph("Test Results", style='CustomHeading')

        # Create table with headers
        table = doc.add_table(rows=1, cols=4)
        table.style = 'Light Grid Accent 1'

        # Header row
        header_cells = table.rows[0].cells
        headers = ["Test Name", "Status", "Value", "Expected"]
        for idx, header in enumerate(headers):
            header_cells[idx].text = header
            header_cells[idx].paragraphs[0].runs[0].font.bold = True
            self._shade_cell(header_cells[idx], RGBColor(31, 71, 136))
            header_cells[idx].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

        # Data rows
        for test in test_results:
            row_cells = table.add_row().cells
            row_cells[0].text = test.get("name", "N/A")
            row_cells[1].text = test.get("status", "N/A")
            row_cells[2].text = str(test.get("value", "N/A"))
            row_cells[3].text = str(test.get("expected", "N/A"))

            # Color code status
            status = test.get("status", "")
            if status == "PASS":
                self._shade_cell(row_cells[1], RGBColor(232, 245, 233))
            elif status == "FAIL":
                self._shade_cell(row_cells[1], RGBColor(255, 235, 238))

        doc.add_paragraph()

    def _add_charts_section(self, doc: 'Document', charts: List[Dict[str, Any]]):
        """
        Add charts and graphs section.

        Args:
            doc: Document instance
            charts: List of chart data
        """
        doc.add_page_break()
        doc.add_paragraph("Charts and Graphs", style='CustomHeading')

        for chart_data in charts:
            # Chart title
            chart_title = chart_data.get("title", "Chart")
            doc.add_paragraph(chart_title, style='Heading 3')

            # Add chart image if available
            if "image_data" in chart_data:
                try:
                    img_data = base64.b64decode(chart_data["image_data"])
                    img_stream = io.BytesIO(img_data)
                    doc.add_picture(img_stream, width=Inches(6))
                except Exception as e:
                    logger.error(f"Failed to add chart image: {e}")
                    doc.add_paragraph(f"[Chart could not be rendered: {e}]")

            # Chart description
            if "description" in chart_data:
                para = doc.add_paragraph(chart_data["description"], style='CustomBody')
                para.paragraph_format.space_after = Pt(12)

            doc.add_paragraph()

    def _add_analysis_section(self, doc: 'Document', analysis: str):
        """
        Add analysis section.

        Args:
            doc: Document instance
            analysis: Analysis text
        """
        doc.add_paragraph("Analysis", style='CustomHeading')
        doc.add_paragraph(analysis, style='CustomBody')
        doc.add_paragraph()

    def _add_recommendations_section(self, doc: 'Document', recommendations: List[str]):
        """
        Add recommendations section.

        Args:
            doc: Document instance
            recommendations: List of recommendations
        """
        doc.add_paragraph("Recommendations", style='CustomHeading')

        for idx, recommendation in enumerate(recommendations, 1):
            para = doc.add_paragraph(style='CustomBody')
            para.add_run(f"{idx}. ").bold = True
            para.add_run(recommendation)

        doc.add_paragraph()

    def _add_appendix_section(self, doc: 'Document', appendix: Dict[str, Any]):
        """
        Add appendix section.

        Args:
            doc: Document instance
            appendix: Appendix data
        """
        doc.add_page_break()
        doc.add_paragraph("Appendix", style='CustomHeading')

        for key, value in appendix.items():
            doc.add_paragraph(key.replace("_", " ").title(), style='Heading 3')
            doc.add_paragraph(str(value), style='CustomBody')
            doc.add_paragraph()

    def _shade_cell(self, cell, color: RGBColor):
        """
        Shade table cell with color.

        Args:
            cell: Table cell
            color: RGB color
        """
        shading_elm = OxmlElement('w:shd')
        shading_elm.set(qn('w:fill'), f'{color.rgb:06X}')
        cell._element.get_or_add_tcPr().append(shading_elm)

    def add_page_break(self, doc: 'Document'):
        """
        Add page break to document.

        Args:
            doc: Document instance
        """
        doc.add_page_break()

    def add_header(self, doc: 'Document', header_text: str):
        """
        Add header to document.

        Args:
            doc: Document instance
            header_text: Header text
        """
        section = doc.sections[0]
        header = section.header
        header_para = header.paragraphs[0]
        header_para.text = header_text
        header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        header_para.runs[0].font.size = Pt(10)
        header_para.runs[0].font.color.rgb = RGBColor(100, 100, 100)

    def add_footer(self, doc: 'Document', footer_text: str):
        """
        Add footer to document.

        Args:
            doc: Document instance
            footer_text: Footer text
        """
        section = doc.sections[0]
        footer = section.footer
        footer_para = footer.paragraphs[0]
        footer_para.text = footer_text
        footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer_para.runs[0].font.size = Pt(9)
        footer_para.runs[0].font.color.rgb = RGBColor(128, 128, 128)

    def validate_output(self, output_path: Path) -> bool:
        """
        Validate Word output.

        Args:
            output_path: Path to Word file

        Returns:
            True if valid
        """
        if not output_path.exists():
            return False

        if output_path.stat().st_size == 0:
            return False

        try:
            # Try to open document
            doc = Document(str(output_path))
            # Check if document has content
            if len(doc.paragraphs) == 0:
                return False
            return True
        except Exception as e:
            logger.error(f"Word validation failed: {e}")
            return False

    def merge_documents(self, documents: List[Path], output_path: Path) -> Path:
        """
        Merge multiple Word documents.

        Args:
            documents: List of document paths
            output_path: Output merged document path

        Returns:
            Path to merged document
        """
        logger.info(f"Merging {len(documents)} documents")

        merged_doc = Document()

        for idx, doc_path in enumerate(documents):
            doc = Document(str(doc_path))

            # Add page break between documents
            if idx > 0:
                merged_doc.add_page_break()

            # Copy all elements
            for element in doc.element.body:
                merged_doc.element.body.append(element)

        merged_doc.save(str(output_path))
        logger.info(f"Documents merged: {output_path}")

        return output_path

    def convert_to_pdf(self, docx_path: Path, pdf_path: Optional[Path] = None) -> Path:
        """
        Convert Word document to PDF.

        Args:
            docx_path: Input Word document path
            pdf_path: Output PDF path (optional)

        Returns:
            Path to PDF file
        """
        if pdf_path is None:
            pdf_path = docx_path.parent / f"{docx_path.stem}.pdf"

        logger.info(f"Converting Word to PDF: {pdf_path}")

        # Note: Conversion requires external tools like LibreOffice or Microsoft Word
        # This is a placeholder for the actual implementation
        logger.warning("PDF conversion requires LibreOffice or MS Word")

        return pdf_path
